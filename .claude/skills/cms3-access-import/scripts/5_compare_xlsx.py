"""Excel đối chiếu danh mục hoá chất WMS (mã C) với Access CMS 3.0 (mã H)."""
import csv
import json
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

src, out = sys.argv[1], sys.argv[2]
rows = json.load(open(f"{src}/compare.json", encoding="utf-8"))
review = list(csv.reader(open(f"{src}/4_review_categories.csv", encoding="utf-8-sig")))
new_ai = list(csv.reader(open(f"{src}/4_review_new_ai.csv", encoding="utf-8-sig")))

FONT = "Arial"
f_norm = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_title = Font(name=FONT, size=14, bold=True, color="1F4E79")
f_note = Font(name=FONT, size=9, italic=True, color="595959")
fill_head = PatternFill("solid", fgColor="1F4E79")
fill_sub = PatternFill("solid", fgColor="DDEBF7")
fill_wms = PatternFill("solid", fgColor="E2EFDA")
fill_acc = PatternFill("solid", fgColor="FFF2CC")
fill_input = PatternFill("solid", fgColor="FFFF00")
fill_bad = PatternFill("solid", fgColor="FFC7CE")
fill_warn = PatternFill("solid", fgColor="FCE4D6")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()

# ------------------------------------------------------------------ Sheet Đối chiếu C-H
ws = wb.active
ws.title = "Đối chiếu C-H"
SH = "'Đối chiếu C-H'"

# (nhóm, tiêu đề cột, khoá dữ liệu, độ rộng, nguồn) nguồn: A=Access, W=WMS, S=so sánh, I=ô nhập, None
cols = [
    ("Mã", "STT", None, 6, None),
    ("Mã", "Mã WMS (C)", "c", 10, "W"),
    ("Mã", "Mã Access (H)", "h", 10, "A"),
    ("Mã", "Tình trạng nạp", "map", 16, None),
    ("Mã", "Mã H khác gộp cùng mã C", "merged", 16, None),
    ("Mã", "Số trường lệch", "#ndiff", 9, None),
    ("Tên hoá chất", "Access - tên TA", "a_ta", 30, "A"),
    ("Tên hoá chất", "Access - tên TV", "a_tv", 26, "A"),
    ("Tên hoá chất", "WMS - tên hoá chất", "w_name", 26, "W"),
    ("Tên hoá chất", "So sánh tên", "cmp:Tên", 12, "S"),
    ("CAS / hoạt chất", "Access - CAS", "a_cas", 16, "A"),
    ("CAS / hoạt chất", "WMS - CAS hoạt chất", "w_cas", 16, "W"),
    ("CAS / hoạt chất", "WMS - hoạt chất (mã, tên, %)", "w_ai", 30, "W"),
    ("CAS / hoạt chất", "WMS - đơn chất / hỗn hợp", "w_type", 11, "W"),
    ("CAS / hoạt chất", "So sánh CAS", "cmp:CAS", 10, "S"),
    ("Nhà sản xuất", "Access - NSX", "a_nsx", 14, "A"),
    ("Nhà sản xuất", "WMS - NSX (viết tắt)", "w_nsx", 14, "W"),
    ("Nhà sản xuất", "WMS - NSX đầy đủ", "w_nsx_full", 22, "W"),
    ("Nhà sản xuất", "So sánh NSX", "cmp:NSX", 10, "S"),
    ("Loại", "Access - Loại", "a_loai", 10, "A"),
    ("Loại", "WMS - Loại", "w_loai", 10, "W"),
    ("Loại", "So sánh loại", "cmp:Loại", 10, "S"),
    ("Trạng thái", "Access - Active", "a_active", 11, "A"),
    ("Trạng thái", "WMS - sử dụng", "w_status", 11, "W"),
    ("Trạng thái", "So sánh trạng thái", "cmp:Trạng thái", 10, "S"),
    ("Phân loại NĐ 24/2026", "Access - nhóm (cờ PL, 11 = cấm)", "a_groups", 12, "A"),
    ("Phân loại NĐ 24/2026", "WMS - nhóm suy ra", "w_groups", 12, "W"),
    ("Phân loại NĐ 24/2026", "So sánh nhóm", "cmp:Nhóm", 10, "S"),
    ("Cảnh báo an toàn", "Access - cờ Độc/Cháy/Ăn mòn/Oxy hoá/Nổ", "a_warn", 20, "A"),
    ("Cảnh báo an toàn", "WMS - cảnh báo", "w_warn", 20, "W"),
    ("Cảnh báo an toàn", "So sánh cảnh báo", "cmp:Cảnh báo", 10, "S"),
    ("Phòng QC1 (QC-B1)", "Access - ĐVT", "a_QC1_dvt", 8, "A"),
    ("Phòng QC1 (QC-B1)", "WMS - ĐVT", "w_QC1_dvt", 8, "W"),
    ("Phòng QC1 (QC-B1)", "So sánh ĐVT", "cmp:ĐVT QC1", 10, "S"),
    ("Phòng QC1 (QC-B1)", "Access - hạn dùng sau mở (tháng)", "a_QC1_snhd", 9, "A"),
    ("Phòng QC1 (QC-B1)", "WMS - hạn dùng (tháng)", "w_QC1_snhd", 9, "W"),
    ("Phòng QC1 (QC-B1)", "So sánh hạn dùng", "cmp:Hạn dùng QC1", 10, "S"),
    ("Phòng QC1 (QC-B1)", "Access - định khu", "a_QC1_dk", 10, "A"),
    ("Phòng QC1 (QC-B1)", "WMS - định khu", "w_QC1_dk", 18, "W"),
    ("Phòng QC2 (QC-B2)", "Access - ĐVT", "a_QC2_dvt", 8, "A"),
    ("Phòng QC2 (QC-B2)", "WMS - ĐVT", "w_QC2_dvt", 8, "W"),
    ("Phòng QC2 (QC-B2)", "So sánh ĐVT", "cmp:ĐVT QC2", 10, "S"),
    ("Phòng QC2 (QC-B2)", "Access - hạn dùng sau mở (tháng)", "a_QC2_snhd", 9, "A"),
    ("Phòng QC2 (QC-B2)", "WMS - hạn dùng (tháng)", "w_QC2_snhd", 9, "W"),
    ("Phòng QC2 (QC-B2)", "So sánh hạn dùng", "cmp:Hạn dùng QC2", 10, "S"),
    ("Phòng QC2 (QC-B2)", "Access - định khu", "a_QC2_dk", 10, "A"),
    ("Phòng QC2 (QC-B2)", "WMS - định khu", "w_QC2_dk", 18, "W"),
    ("Kết luận (người kiểm tra điền)", "Kết luận / hướng xử lý", "#in", 30, "I"),
    ("Kết luận (người kiểm tra điền)", "Người kiểm tra", "#in", 14, "I"),
    ("Kết luận (người kiểm tra điền)", "Ngày kiểm tra", "#in", 11, "I"),
]
cmp_idx = [i + 1 for i, c in enumerate(cols) if c[4] == "S"]
col_of = {c[1] + "|" + c[0]: i + 1 for i, c in enumerate(cols)}

# Hàng 1: nhóm (merge), hàng 2: tiêu đề cột
start = 1
for i in range(1, len(cols) + 1):
    if i == len(cols) or cols[i][0] != cols[start - 1][0]:
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=i)
        cell = ws.cell(1, start, cols[start - 1][0])
        cell.font, cell.fill, cell.alignment = f_head, fill_head, center
        start = i + 1
for i, (_, title, _, width, srcflag) in enumerate(cols, 1):
    cell = ws.cell(2, i, title)
    cell.font, cell.alignment, cell.border = f_bold, center, border
    cell.fill = {"A": fill_acc, "W": fill_wms, "I": fill_input}.get(srcflag, fill_sub)
    ws.column_dimensions[get_column_letter(i)].width = width
ws.row_dimensions[2].height = 45

cmp_first, cmp_last = get_column_letter(cmp_idx[0]), get_column_letter(cmp_idx[-1])
r0 = 3
for n, row in enumerate(sorted(rows, key=lambda x: (x["c"] == "", x["c"], x["h"])), 0):
    rr = r0 + n
    cmp = row.get("cmp") or {}
    for i, (_, title, key, _, srcflag) in enumerate(cols, 1):
        if title == "STT":
            v = n + 1
        elif key == "#ndiff":
            # Đếm các ô so sánh "Khác" / "Thiếu ở WMS" trên cùng dòng
            v = f'=COUNTIF({cmp_first}{rr}:{cmp_last}{rr},"Khác")+COUNTIF({cmp_first}{rr}:{cmp_last}{rr},"Thiếu ở WMS")'
        elif key == "#in":
            v = None
        elif key.startswith("cmp:"):
            v = cmp.get(key[4:], "")
        else:
            v = row.get(key, "")
        c = ws.cell(rr, i, v if v != "" else None)
        c.font, c.border = f_norm, border
        c.alignment = center if srcflag == "S" or key == "#ndiff" or title == "STT" else wrap
        if srcflag == "I":
            c.fill = fill_input
last = r0 + len(rows) - 1

# Định dạng có điều kiện cho các cột so sánh
for i in cmp_idx:
    rng = f"{get_column_letter(i)}{r0}:{get_column_letter(i)}{last}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Khác"'], fill=fill_bad, font=Font(name=FONT, bold=True, color="9C0006")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Thiếu ở WMS"'], fill=fill_warn, font=Font(name=FONT, bold=True, color="C65911")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Đúng"'], font=Font(name=FONT, color="548235")))
ndiff_col = get_column_letter(6)
ws.conditional_formatting.add(f"{ndiff_col}{r0}:{ndiff_col}{last}",
                              CellIsRule(operator="greaterThan", formula=["0"], fill=fill_bad, font=Font(name=FONT, bold=True, color="9C0006")))
ws.freeze_panes = "D3"
ws.auto_filter.ref = f"A2:{get_column_letter(len(cols))}{last}"
dv = DataValidation(type="list", formula1='"Đúng - giữ nguyên,Đã sửa trên WMS,Cần sửa trên WMS,Access sai - bỏ qua"', allow_blank=True)
dv.error, dv.showErrorMessage = "Chọn trong danh sách hoặc gõ nội dung khác", False
ws.add_data_validation(dv)
kl_col = get_column_letter(len(cols) - 2)
dv.add(f"{kl_col}{r0}:{kl_col}{last}")

# ------------------------------------------------------------------ Sheet Cần kiểm tra
ws2 = wb.create_sheet("Mã nghi ngờ")
head2 = ["STT"] + review[0]
widths2 = [6, 10, 10, 14, 26, 30, 12, 8, 11, 28, 9, 9, 60, 30, 14, 11]
for i, h in enumerate(head2, 1):
    c = ws2.cell(1, i, h)
    c.font, c.fill, c.alignment, c.border = f_head, fill_head, center, border
    ws2.column_dimensions[get_column_letter(i)].width = widths2[i - 1] if i <= len(widths2) else 14
ws2.row_dimensions[1].height = 30
for n, rec in enumerate(review[1:], 0):
    rr = n + 2
    for i, v in enumerate([n + 1] + rec, 1):
        c = ws2.cell(rr, i, v if v != "" else None)
        c.font, c.border, c.alignment = f_norm, border, wrap
        if i >= len(head2) - 2:
            c.fill = fill_input
last2 = len(review)
rng_p = f"B2:B{last2}"
ws2.conditional_formatting.add(rng_p, CellIsRule(operator="equal", formula=['"Cao"'], fill=fill_bad, font=Font(name=FONT, bold=True, color="9C0006")))
ws2.conditional_formatting.add(rng_p, CellIsRule(operator="equal", formula=['"Trung bình"'], fill=fill_warn))
ws2.freeze_panes = "E2"
ws2.auto_filter.ref = f"A1:{get_column_letter(len(head2))}{last2}"
dv2 = DataValidation(type="list", formula1='"Đúng - giữ nguyên,Đã sửa trên WMS,Cần sửa trên WMS,Access sai - bỏ qua"', allow_blank=True)
dv2.showErrorMessage = False
ws2.add_data_validation(dv2)
dv2.add(f"{get_column_letter(len(head2) - 2)}2:{get_column_letter(len(head2) - 2)}{last2}")

# ------------------------------------------------------------------ Sheet Hoạt chất tạo mới
ws3 = wb.create_sheet("Hoạt chất tạo mới")
widths3 = [9, 28, 28, 13, 14, 9, 22, 40, 26]
for i, h in enumerate(new_ai[0], 1):
    c = ws3.cell(1, i, h)
    c.font, c.fill, c.alignment, c.border = f_head, fill_head, center, border
    ws3.column_dimensions[get_column_letter(i)].width = widths3[i - 1]
for n, rec in enumerate(new_ai[1:], 2):
    for i, v in enumerate(rec, 1):
        c = ws3.cell(n, i, v if v != "" else None)
        c.font, c.border, c.alignment = f_norm, border, wrap
        if i == len(new_ai[0]):
            c.fill = fill_input
ws3.freeze_panes = "C2"
ws3.auto_filter.ref = f"A1:{get_column_letter(len(new_ai[0]))}{len(new_ai)}"

# ------------------------------------------------------------------ Sheet Tổng hợp (đặt đầu)
ws0 = wb.create_sheet("Tổng hợp", 0)
ws0.column_dimensions["A"].width = 46
ws0.column_dimensions["B"].width = 12
ws0.column_dimensions["C"].width = 12
ws0.column_dimensions["D"].width = 12
ws0.column_dimensions["E"].width = 70
ws0["A1"] = "ĐỐI CHIẾU DANH MỤC HOÁ CHẤT WMS (MÃ C) VỚI ACCESS CMS 3.0 (MÃ H)"
ws0["A1"].font = f_title
ws0["A2"] = (f"Nguồn: Access Z:\\CMS_3.0_L.accdb (Q_DMHC_ChungHH, danh mục phòng QC1/QC2) xuất ngày {date.today():%d/%m/%Y}; "
             "WMS: bảng chemical_categories / chemical_department_categories cùng ngày. Mã H ↔ mã C lấy từ lịch sử 'Thêm mới' của danh mục.")
ws0["A2"].font = f_note
ws0.merge_cells("A2:E2")
ws0["A2"].alignment = wrap
ws0.row_dimensions[2].height = 30

r = 4
for i, h in enumerate(["Tiêu chí so sánh", "Đúng", "Khác", "Thiếu ở WMS", "Ghi chú"], 1):
    c = ws0.cell(r, i, h)
    c.font, c.fill, c.alignment, c.border = f_head, fill_head, center, border
notes = {
    "So sánh tên": "'Theo tên hoạt chất' = đơn chất đặt tên theo hoạt chất (quy tắc đã chốt), không tính là lệch",
    "So sánh CAS": "So CAS Access với CAS các hoạt chất liên kết; 'Nước' tự thêm cho dung dịch không tính",
    "So sánh NSX": "Access NSX với tên viết tắt NSX trên WMS",
    "So sánh trạng thái": "Nhiều mã H gộp 1 mã C: mã C hoạt động nếu có ít nhất 1 mã H hoạt động",
    "So sánh nhóm": "Nhóm WMS tự suy từ hoạt chất (ChemicalClassification); 11 = hoá chất cấm",
    "So sánh cảnh báo": "Cảnh báo WMS suy từ cờ Access lúc nạp (Ăn mòn: axit/bazơ đoán theo tên)",
    "So sánh ĐVT": "Mã gộp có ĐVT khác nhau thì WMS giữ 1 ĐVT",
    "So sánh hạn dùng": "SNHD Access (tháng sau mở nắp) với hạn dùng của danh mục phòng",
}
for i in cmp_idx:
    r += 1
    group, title = cols[i - 1][0], cols[i - 1][1]
    L = get_column_letter(i)
    rng = f"{SH}!{L}${r0}:{L}${last}"
    label = title if not group.startswith("Phòng") else f"{title} - {group}"
    ws0.cell(r, 1, label)
    ws0.cell(r, 2, f'=COUNTIF({rng},"Đúng")+COUNTIF({rng},"Theo tên hoạt chất")')
    ws0.cell(r, 3, f'=COUNTIF({rng},"Khác")')
    ws0.cell(r, 4, f'=COUNTIF({rng},"Thiếu ở WMS")')
    ws0.cell(r, 5, notes.get(title, ""))
    for k in range(1, 6):
        c = ws0.cell(r, k)
        c.font, c.border = f_norm, border
        c.alignment = center if 2 <= k <= 4 else wrap
        if 2 <= k <= 4:
            c.number_format = "#,##0"
summary_last = r

r += 2
for i, h in enumerate(["Tổng quan", "Số lượng", "", "", "Ghi chú"], 1):
    c = ws0.cell(r, i, h if h else None)
    c.font, c.fill, c.alignment, c.border = f_head, fill_head, center, border
D = get_column_letter(4)
items = [
    ("Số mã H trong Access (Q_DMHC_ChungHH, đã bỏ trùng)", f'=COUNTA({SH}!C{r0}:C{last})', "Mỗi mã H một dòng ở sheet 'Đối chiếu C-H'"),
    ("  - Đã nạp vào WMS", f'=COUNTIF({SH}!{D}{r0}:{D}{last},"Đã nạp")', ""),
    ("  - Không nạp (dữ liệu thử OQ..)", f'=COUNTIF({SH}!{D}{r0}:{D}{last},"Không nạp - dữ liệu thử*")', "Tên OQ1/OQ02, Loại OQLHC.., ĐVT OQDV.."),
    ("Số mã C trên WMS có mã Access", f'=SUMPRODUCT(({SH}!B{r0}:B{last}<>"")/COUNTIF({SH}!B{r0}:B{last},{SH}!B{r0}:B{last}&""))', "Nhiều mã H gộp 1 mã C (cùng tên + NSX + Loại)"),
    ("Số mã H có ít nhất 1 trường lệch", f'=COUNTIF({SH}!F{r0}:F{last},">0")', "Lọc cột 'Số trường lệch' > 0 ở sheet 'Đối chiếu C-H'"),
    ("Số dòng đã có kết luận (Đối chiếu C-H)", f'=COUNTA({SH}!{kl_col}{r0}:{kl_col}{last})', ""),
    ("Mã nghi ngờ - mức Cao", f"=COUNTIF('Mã nghi ngờ'!B2:B{last2},\"Cao\")", "Sheet 'Mã nghi ngờ' (theo quy tắc rà soát bước 4)"),
    ("Mã nghi ngờ - mức Trung bình", f"=COUNTIF('Mã nghi ngờ'!B2:B{last2},\"Trung bình\")", ""),
    ("Mã nghi ngờ - mức Thấp", f"=COUNTIF('Mã nghi ngờ'!B2:B{last2},\"Thấp\")", ""),
    ("Mã nghi ngờ đã có kết luận", f"=COUNTA('Mã nghi ngờ'!{get_column_letter(len(head2) - 2)}2:{get_column_letter(len(head2) - 2)}{last2})", ""),
    ("Hoạt chất tạo mới (ngoài danh mục luật)", f"=COUNTA('Hoạt chất tạo mới'!A2:A{len(new_ai)})", "Rà tên / CAS / nhóm"),
]
for label, formula, note in items:
    r += 1
    ws0.cell(r, 1, label)
    ws0.cell(r, 2, formula)
    ws0.cell(r, 5, note)
    for k in (1, 2, 5):
        c = ws0.cell(r, k)
        c.font, c.border = f_norm, border
        c.alignment = center if k == 2 else wrap
    ws0.cell(r, 2).number_format = "#,##0"

r += 2
for i, h in enumerate(["Vấn đề (sheet 'Mã nghi ngờ')", "Số mã", "", "", "Cách xử lý gợi ý"], 1):
    c = ws0.cell(r, i, h if h else None)
    c.font, c.fill, c.alignment, c.border = f_head, fill_head, center, border
VC = get_column_letter(head2.index("Vấn đề") + 1)
issues = [
    ("Lệch nhóm Access ≠ WMS", "*Lệch nhóm*", "Kiểm tra hoạt chất / % thành phần của tên hoá chất"),
    ("Access nhóm 10 nhưng WMS không suy ra", "*Access ghi nhóm 10*", "Thêm thành phần nhóm 9 + tick Bảng B"),
    ("Bảng B đang tick mặc định", "*Bảng B đang tick*", "Chọn đúng mục nguy hại ở Dữ liệu gốc → Hoá chất"),
    ("Hoạt chất tạo mới tự gắn mục gộp", "*tự gắn mục gộp*", "Kiểm tra mục gộp, hệ số quy đổi"),
    ("Hoạt chất ngoài danh mục luật được gán nhóm", "*ngoài danh mục luật*", "Xác nhận nhóm theo Access có đúng"),
    ("Thành phần nhóm 3-7 chưa có %", "*chưa có %*", "Khai % để xét nhóm 8"),
    ("Dung dịch pha loãng (tự thêm Nước)", "*pha loãng*", "Xác nhận"),
    ("Đơn chất mất nồng độ / cấp độ trong tên", "*mất nồng độ*", "Cấp độ còn ở cột Loại + NSX; đổi tên nếu cần"),
    ("Đơn chất không có CAS", "*không có số CAS*", "Bổ sung CAS cho hoạt chất"),
    ("Tên dạng mã / công thức", "*dạng mã/công thức*", "Đặt lại tên đầy đủ"),
    ("Loại lạ", "*Loại lạ*", "Sửa cột Loại"),
    ("Tên trùng nên bị gắn mã Access", "*gắn mã Access*", "Đặt lại tên"),
    ("Mã đang khoá", "*đang khoá*", "Xác nhận còn dùng hay không"),
    ("Tương kỵ cùng tầng định khu", "*Tương kỵ*", "Xếp lại định khu"),
]
for label, pattern, how in issues:
    r += 1
    ws0.cell(r, 1, label)
    ws0.cell(r, 2, f"=COUNTIF('Mã nghi ngờ'!{VC}2:{VC}{last2},\"{pattern}\")")
    ws0.cell(r, 5, how)
    for k in (1, 2, 5):
        c = ws0.cell(r, k)
        c.font, c.border = f_norm, border
        c.alignment = center if k == 2 else wrap
    ws0.cell(r, 2).number_format = "#,##0"

r += 2
ws0.cell(r, 1, "CÁCH DÙNG").font = f_bold
legend = [
    ("Ô nền vàng", "Người kiểm tra điền: Kết luận / hướng xử lý, Người kiểm tra, Ngày kiểm tra (chọn sẵn hoặc gõ tự do)."),
    ("Ví dụ kết luận", "'Cần sửa trên WMS: NSX đổi thành Merck' ; 'Access sai - bỏ qua' ; 'Đúng - giữ nguyên'."),
    ("Cột nền vàng nhạt / xanh nhạt", "Tiêu đề vàng nhạt = dữ liệu Access, xanh nhạt = dữ liệu WMS, xanh dương = kết quả so sánh."),
    ("Lọc nhanh", "Sheet 'Đối chiếu C-H': lọc 'Số trường lệch' > 0, hoặc lọc từng cột 'So sánh ...' = Khác."),
    ("Mã H gộp", "Nhiều mã H cùng tên + NSX + Loại được gộp vào 1 mã C; cột 'Mã H khác gộp cùng mã C' liệt kê các mã còn lại."),
]
for a, b in legend:
    r += 1
    ws0.cell(r, 1, a).font = f_bold
    ws0.cell(r, 2, b).font = f_norm
    ws0.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    ws0.cell(r, 2).alignment = wrap
    ws0.row_dimensions[r].height = 28
ws0.cell(r - 4, 1).fill = fill_input

wb.calculation.fullCalcOnLoad = True
wb.save(out)
print("saved", out, "rows", len(rows), "review", len(review) - 1, "ai", len(new_ai) - 1)
